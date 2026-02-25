import os
import shutil
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# Paths
# ─────────────────────────────────────────────
BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = (
    BASE_DIR
    / "input"
    / "UPLOAD PM"
    / "QTKD (OK)"
    / "Lễ trao chứng chỉ CTH cho sv QTKS 05.12.2025"
)
OUTPUT_DIR = BASE_DIR / "5.2" / "excel" / "Lễ trao chứng chỉ CTH cho sv QTKS 05.12.2025"

# ─────────────────────────────────────────────
# Styles
# ─────────────────────────────────────────────
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", start_color="2E74B5")
CELL_FONT = Font(name="Arial", size=10)
ALT_FILL = PatternFill("solid", start_color="EBF3FB")
BORDER_SIDE = Side(style="thin", color="BFBFBF")
CELL_BORDER = Border(
    left=BORDER_SIDE, right=BORDER_SIDE, top=BORDER_SIDE, bottom=BORDER_SIDE
)
TEXT_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
TEXT_HEADER_FILL = PatternFill("solid", start_color="404040")


def auto_width(ws, min_width=10, max_width=60):
    for col in ws.columns:
        length = max(
            len(str(cell.value)) if cell.value is not None else 0 for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(length + 2, min_width), max_width
        )


def write_text_sheet(ws, paragraphs: list[str]):
    # Header row
    ws.append(["#", "Paragraph"])
    for cell in ws[1]:
        cell.font = TEXT_HEADER_FONT
        cell.fill = TEXT_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 80
    ws.row_dimensions[1].height = 20

    for i, text in enumerate(paragraphs, start=1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=text)
        fill = ALT_FILL if i % 2 == 0 else PatternFill()
        for col in (1, 2):
            c = ws.cell(row=row, column=col)
            c.font = CELL_FONT
            c.border = CELL_BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if fill.fill_type:
                c.fill = fill

    ws.freeze_panes = "A2"


def write_table_sheet(ws, table_data: list[list[str]], table_index: int):
    if not table_data:
        return

    # First row → header
    headers = table_data[0]
    ws.append(headers)
    for col_idx, cell in enumerate(ws[1], start=1):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = CELL_BORDER
    ws.row_dimensions[1].height = 22

    for row_idx, row in enumerate(table_data[1:], start=2):
        ws.append(row)
        fill = ALT_FILL if row_idx % 2 == 0 else PatternFill()
        for col_idx, cell in enumerate(ws[row_idx], start=1):
            cell.font = CELL_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill.fill_type:
                cell.fill = fill

    auto_width(ws)
    ws.freeze_panes = "A2"


def docx_to_xlsx(docx_path: Path, output_path: Path):
    doc = Document(str(docx_path))
    wb = Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    # ── Extract paragraphs (non-empty) ──
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    if paragraphs:
        ws_text = wb.create_sheet("Text")
        write_text_sheet(ws_text, paragraphs)

    # ── Extract tables ──
    for i, table in enumerate(doc.tables, start=1):
        sheet_name = f"Table_{i}"
        ws = wb.create_sheet(sheet_name)
        table_data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        write_table_sheet(ws, table_data, i)

    if not wb.sheetnames:
        ws = wb.create_sheet("Empty")
        ws["A1"] = "No content found in document."

    wb.save(str(output_path))


def collect_word_files() -> list[Path]:
    files = []
    for root, dirs, filenames in os.walk(INPUT_DIR):
        dirs[:] = [
            d
            for d in dirs
            if not d.startswith(".") and d not in ("venv", "__pycache__")
        ]
        for f in filenames:
            if f.lower().endswith((".docx", ".doc")) and not f.startswith("output"):
                files.append(Path(root) / f)
    return sorted(files)


def run():
    if not INPUT_DIR.exists():
        print(f"❌ Không tìm thấy thư mục: {INPUT_DIR}")
        return

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    word_files = collect_word_files()
    if not word_files:
        print("⚠️  Không tìm thấy file Word nào.")
        return

    print(f"📂 Quét từ  : {INPUT_DIR}")
    print(f"📁 Lưu vào  : {OUTPUT_DIR}")
    print(f"📝 Tìm thấy : {len(word_files)} file")
    print("-" * 50)

    success, failed = 0, []

    for src in word_files:
        out_name = src.stem + ".xlsx"
        out_path = OUTPUT_DIR / out_name

        # Handle collisions
        counter = 1
        while out_path.exists():
            out_path = OUTPUT_DIR / f"{src.stem}_{counter}.xlsx"
            counter += 1

        try:
            docx_to_xlsx(src, out_path)
            rel = src.relative_to(INPUT_DIR)
            print(f"  ✅ {rel}  →  {out_path.name}")
            success += 1
        except Exception as e:
            print(f"  ❌ {src.name}: {e}")
            failed.append(src.name)

    print("-" * 50)
    print(f"✅ Thành công: {success} file")
    if failed:
        print(f"❌ Thất bại  : {len(failed)} file")
        for f in failed:
            print(f"   - {f}")


if __name__ == "__main__":
    run()
